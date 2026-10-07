#!/usr/bin/env python3
"""견적서 생성: python3 make_quote.py spec.json OUT_DIR  -> quote.html(Artifact용), 견적서_*.xlsx, 견적서_*.pdf
spec: recipient, date, items[{name,qty,unit,price}], qty_suffix, discount_pct, discount_label,
      vat("exempt"|"10"), schedule, account, notes[], supplier{name,regno,ceo,phone}"""
import json, os, subprocess, sys
D = os.path.dirname(os.path.abspath(__file__))
spec = json.load(open(sys.argv[1])); out = sys.argv[2]; os.makedirs(out, exist_ok=True)
sup = {'name': '인트린직 (서퍼스)', 'regno': '146-33-01415', 'ceo': '김승렬', 'phone': '010-9016-1681', **spec.get('supplier', {})}
won = lambda n: f"{n:,}원"
sfx = spec.get('qty_suffix', '명')
sub = sum(i['qty'] * i['price'] for i in spec['items'])
pct = spec.get('discount_pct', 0)
disc = round(sub * pct / 100)
net = sub - disc
vat = 0 if spec.get('vat', 'exempt') == 'exempt' else round(net * 0.1)
total = net + vat

rows = ''.join(f'<tr><td class="item-name">{i["name"]}</td><td class="qty">{i["qty"]}{sfx}</td>'
               f'<td class="price">{won(i["price"])}</td><td class="amount">{won(i["qty"]*i["price"])}</td></tr>' for i in spec['items'])
r = lambda cls, label, v: f'<tr class="{cls}"><td colspan="3" class="label">{label}</td><td>{v}</td></tr>'
foot = ''
if pct:
    foot += r('subtotal', '소계', won(sub)) + r('subtotal', spec.get('discount_label', f'할인 ({pct}%)') , '-' + won(disc))
foot += r('subtotal', '합계', won(net))
foot += r('vat', '부가세 (면제)' if vat == 0 and spec.get('vat', 'exempt') == 'exempt' else '부가세 (10%)', won(vat))
foot += r('total', '총액', won(total))
info = [('교육 일정', spec.get('schedule')), ('입금 계좌', spec.get('account'))]
info = ''.join(f'<tr><th>{k}</th><td>{v}</td></tr>' for k, v in info if v)
info = f'<table class="supplier info">{info}</table>' if info else ''
notes = '<br>'.join('※ ' + n for n in spec.get('notes', []))

t = open(f'{D}/template.html').read()
for k, v in {'{{RECIPIENT}}': spec['recipient'], '{{DATE}}': spec['date'], '{{SUPPLIER_NAME}}': sup['name'],
             '{{SUPPLIER_REGNO}}': sup['regno'], '{{SUPPLIER_CEO}}': sup['ceo'], '{{SUPPLIER_PHONE}}': sup['phone'],
             '{{ITEM_ROWS}}': rows, '{{FOOT_ROWS}}': foot, '{{INFO_TABLE}}': info, '{{NOTES}}': notes,
             '{{SIGNATURE_SVG}}': open(f'{D}/signature.svg').read()}.items():
    t = t.replace(k, v)
assert '{{' not in t
open(f'{out}/quote.html', 'w').write(t)
base = f'견적서_{spec["recipient"]}_{spec["date"].replace("년 ","").replace("월 ","").replace("일","")}'

# PDF: 로컬 렌더용으로 doctype/charset 을 씌우고 A4 한 장으로 인쇄 (Artifact 파일 자체는 건드리지 않음)
p = t.replace('</style>', '@page{size:A4;margin:0}html,body{background:#fff!important;padding:0}body{padding:30px 0!important}.sheet{box-shadow:none!important;margin:0 auto}</style>', 1)
open(f'{out}/print.html', 'w').write('<!doctype html><html lang="ko"><head><meta charset="utf-8">' + p + '</head></html>')
subprocess.run(['/Applications/Google Chrome.app/Contents/MacOS/Google Chrome', '--headless', '--disable-gpu',
                '--no-pdf-header-footer', f'--print-to-pdf={out}/{base}.pdf', f'{out}/print.html'], capture_output=True)

# XLSX (수식 사용 — 이후 xlsx 스킬의 recalc.py 로 값 계산)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
wb = Workbook(); ws = wb.active; ws.title = '견적서'
F = lambda **k: Font(name='Arial', **k)
ws['A1'] = '견 적 서'; ws.merge_cells('A1:D1'); ws['A1'].alignment = Alignment(horizontal='center')
ws['A3'] = '받는자'; ws['B3'] = spec['recipient'] + ' 귀하'; ws['A4'] = '작성일자'; ws['B4'] = spec['date']
for n, (k, v) in enumerate([('상호', sup['name']), ('사업자등록번호', sup['regno']), ('대표자', sup['ceo']), ('연락처', sup['phone'])], 6):
    ws.cell(n, 1, k); ws.cell(n, 2, v)
for c, h in enumerate(['품목', '수량', '단가', '금액'], 1):
    x = ws.cell(11, c, h); x.alignment = Alignment(horizontal='center'); x.border = Border(bottom=Side(style='thin')); x.font = F(bold=True)
n = 12
for i in spec['items']:
    ws.cell(n, 1, i['name']); ws.cell(n, 2, i['qty']); ws.cell(n, 3, i['price']); ws.cell(n, 4, f'=B{n}*C{n}'); n += 1
last = n - 1; n += 1
def line(label, f, bold=False):
    global n
    ws.cell(n, 3, label).font = F(bold=bold); ws.cell(n, 4, f).font = F(bold=bold); n += 1; return n - 1
s = line('소계', f'=SUM(D12:D{last})'); cur = s
if pct:
    d = line(spec.get('discount_label', f'할인 ({pct}%)'), f'=-ROUND(D{s}*{pct}/100,0)'); cur = line('합계', f'=D{s}+D{d}')
else:
    ws.cell(s, 3, '합계')
v = line('부가세 (면제)', 0) if spec.get('vat', 'exempt') == 'exempt' else line('부가세 (10%)', f'=ROUND(D{cur}*0.1,0)')
line('총액', f'=D{cur}+D{v}', True); n += 1
for k, val in [('교육 일정', spec.get('schedule')), ('입금 계좌', spec.get('account'))]:
    if val: ws.cell(n, 1, k); ws.cell(n, 2, val); n += 1
n += 1
for t_ in spec.get('notes', []): ws.cell(n, 1, '※ ' + t_); n += 1
for row in ws.iter_rows():
    for c in row:
        if c.value is not None and c.font.name != 'Arial': c.font = F(bold=c.font.bold, size=18 if c.coordinate == 'A1' else 11)
        if c.column in (3, 4) and c.row >= 12: c.number_format = '#,##0;-#,##0'; c.alignment = Alignment(horizontal='right')
for col, w in zip('ABCD', (24, 22, 22, 16)): ws.column_dimensions[col].width = w
wb.save(f'{out}/{base}.xlsx')
print(json.dumps({'html': f'{out}/quote.html', 'pdf': f'{out}/{base}.pdf', 'xlsx': f'{out}/{base}.xlsx', 'sub': sub, 'discount': disc, 'vat': vat, 'total': total}, ensure_ascii=False))
