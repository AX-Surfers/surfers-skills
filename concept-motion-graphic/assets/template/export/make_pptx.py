"""clips/ 의 스텝별 MP4 로 PPTX 를 만든다.
- 스텝 1개 = 슬라이드 1장, 영상이 화면을 꽉 채우고 슬라이드가 열리면 자동 재생된다.
- 영상이 끝나면 마지막 프레임에서 멈춰 있다 → 클리커로 '다음'을 누르면 다음 스텝.
- 포스터(재생 전 정지 화면)는 각 스텝의 첫 프레임 = 이전 스텝의 끝 화면이라 넘어갈 때 튀지 않는다.
사용법: python3 make_pptx.py [--v basic]  →  <폴더명>.pptx
"""
import json
import os
import sys
from copy import deepcopy

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
SUF = "_basic" if "--v" in sys.argv and sys.argv[sys.argv.index("--v") + 1] == "basic" else ""  # --v basic → 초보자 버전
CLIPS = os.path.join(HERE, "clips" + SUF)
PROJ = os.path.basename(os.path.dirname(HERE))
OUT = os.path.join(HERE, f"{PROJ}{SUF}.pptx")
NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

TIMING = """<p:timing xmlns:p="{ns}">
 <p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
  <p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
   <p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
    <p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
     <p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
      <p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{ms}" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd>
     </p:childTnLst></p:cTn></p:par>
    </p:childTnLst></p:cTn></p:par>
   </p:childTnLst></p:cTn></p:par>
  </p:childTnLst></p:cTn>
  <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
  <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
  </p:seq>
  <p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
 </p:childTnLst></p:cTn></p:par></p:tnLst>
</p:timing>"""


def main():
    steps = json.load(open(os.path.join(CLIPS, "manifest.json"), encoding="utf-8"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    blank = prs.slide_layouts[6]
    for s in steps:
        base = os.path.splitext(s["file"])[0]
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        from pptx.dml.color import RGBColor
        bg.fore_color.rgb = RGBColor(0, 0, 0)
        mv = slide.shapes.add_movie(
            os.path.join(CLIPS, s["file"]), 0, 0, prs.slide_width, prs.slide_height,
            poster_frame_image=os.path.join(CLIPS, base + "_시작.png"), mime_type="video/mp4",
        )
        # python-pptx 가 넣은 '클릭하면 재생' 타이밍을 '자동 재생' 으로 교체
        sld = slide._element
        for t in sld.findall(f"{{{NS}}}timing"):
            sld.remove(t)
        timing = etree.fromstring(TIMING.format(ns=NS, ms=int(s["seconds"] * 1000), spid=mv.shape_id))
        ext = sld.find(f"{{{NS}}}extLst")
        if ext is not None:
            ext.addprevious(timing)
        else:
            sld.append(timing)
        slide.notes_slide.notes_text_frame.text = f'{s["g"]}. {s["title"]} — {s["note"]}'
    prs.save(OUT)
    print(OUT, f"{os.path.getsize(OUT) / 1e6:.1f} MB", len(steps), "slides")


if __name__ == "__main__":
    main()
